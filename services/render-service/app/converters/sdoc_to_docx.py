"""
Convierte el contenido JSON de un documento .sdoc a DOCX usando python-docx.
"""
import json
import logging
from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

logger = logging.getLogger(__name__)

ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _set_run_marks(run, marks: list[dict[str, Any]]) -> None:
    for mark in marks:
        mark_type = mark.get("type", "")
        attrs = mark.get("attrs", {})
        if mark_type == "bold":
            run.bold = True
        elif mark_type == "italic":
            run.italic = True
        elif mark_type == "underline":
            run.underline = True
        elif mark_type == "strike":
            run.font.strike = True
        elif mark_type == "code":
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        elif mark_type == "textStyle":
            color = attrs.get("color", "")
            if color and color.startswith("#"):
                hex_color = color.lstrip("#")
                if len(hex_color) == 6:
                    r, g, b = int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
                    run.font.color.rgb = RGBColor(r, g, b)


def _process_inline(paragraph, node: dict[str, Any]) -> None:
    node_type = node.get("type", "")
    if node_type == "text":
        run = paragraph.add_run(node.get("text", ""))
        _set_run_marks(run, node.get("marks", []))
    elif node_type == "hardBreak":
        paragraph.add_run().add_break()
    elif node_type == "image":
        pass  # Las imágenes requieren descarga previa; se omiten en DOCX básico


def _process_block(doc: Document, node: dict[str, Any], list_level: int = 0) -> None:
    node_type = node.get("type", "")
    content = node.get("content", [])
    attrs = node.get("attrs", {})

    if node_type == "paragraph":
        para = doc.add_paragraph()
        para.alignment = ALIGN_MAP.get(attrs.get("textAlign", "left"), WD_ALIGN_PARAGRAPH.LEFT)
        for child in content:
            _process_inline(para, child)

    elif node_type == "heading":
        level = min(attrs.get("level", 1), 6)
        para = doc.add_heading(level=level)
        para.alignment = ALIGN_MAP.get(attrs.get("textAlign", "left"), WD_ALIGN_PARAGRAPH.LEFT)
        for child in content:
            _process_inline(para, child)

    elif node_type in ("bulletList", "orderedList"):
        for item in content:
            _process_list_item(doc, item, node_type, list_level)

    elif node_type == "blockquote":
        for child in content:
            para = doc.add_paragraph(style="Quote")
            if child.get("type") == "paragraph":
                for inline in child.get("content", []):
                    _process_inline(para, inline)

    elif node_type == "codeBlock":
        text = "".join(
            c.get("text", "") for c in content if c.get("type") == "text"
        )
        para = doc.add_paragraph(text, style="No Spacing")
        para.runs[0].font.name = "Courier New" if para.runs else None

    elif node_type == "horizontalRule":
        para = doc.add_paragraph()
        pPr = para._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "999999")
        pBdr.append(bottom)
        pPr.append(pBdr)

    elif node_type == "table":
        _process_table(doc, node)

    elif node_type == "doc":
        for child in content:
            _process_block(doc, child)


def _process_list_item(doc: Document, item: dict[str, Any], list_type: str, level: int) -> None:
    style = "List Bullet" if list_type == "bulletList" else "List Number"
    if level > 0:
        style += f" {min(level + 1, 3)}"
    for child in item.get("content", []):
        if child.get("type") == "paragraph":
            para = doc.add_paragraph(style=style)
            for inline in child.get("content", []):
                _process_inline(para, inline)
        elif child.get("type") in ("bulletList", "orderedList"):
            for sub_item in child.get("content", []):
                _process_list_item(doc, sub_item, child["type"], level + 1)


def _process_table(doc: Document, table_node: dict[str, Any]) -> None:
    rows = table_node.get("content", [])
    if not rows:
        return
    col_count = max(len(row.get("content", [])) for row in rows)
    table = doc.add_table(rows=0, cols=col_count)
    table.style = "Table Grid"

    for row_node in rows:
        row = table.add_row()
        for idx, cell_node in enumerate(row_node.get("content", [])):
            if idx >= col_count:
                break
            cell = row.cells[idx]
            cell.text = ""
            for child in cell_node.get("content", []):
                if child.get("type") == "paragraph":
                    para = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
                    for inline in child.get("content", []):
                        _process_inline(para, inline)


def convert(content_json: Any, title: str = "Documento") -> bytes:
    """Punto de entrada: retorna bytes del archivo DOCX."""
    if isinstance(content_json, str):
        content_json = json.loads(content_json)

    doc = Document()

    # Configurar márgenes A4
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)

    # Título del documento
    doc.add_heading(title, level=0)

    _process_block(doc, content_json)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
